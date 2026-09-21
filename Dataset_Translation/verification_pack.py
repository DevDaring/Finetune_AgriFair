"""Build the human-verification package for the Hindi and Bengali translations.

    python -m Dataset_Translation.verification_pack [--facts 100] [--pairs 50] [--seed 20260921]

Per language: a stratified sample (facts by axis x condition, pairs by toggle axis), one
Excel workbook per rater with the English source next to the translation and empty rating
columns, a one-page instruction sheet, the glossary, and a manifest with the item ids and
their SHA-256 so returned sheets can be validated. Output: Submission2/verification/*.zip.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import random
import zipfile
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from Dataset_Translation.glossary import GLOSSARY
from Dataset_Translation.run import CFG, OUT, _jsonl

ROOT = Path(__file__).resolve().parents[2]
DEST = ROOT / "Submission2" / "verification"
LANG_NAME = {"hi": "Hindi", "bn": "Bengali"}
RATERS = ["R1", "R2"]

FACT_COLS = ["item_id", "axis", "condition", "english_question", "english_options", "translated_question", "translated_options",
             "meaning_preserved (yes / partly / no)", "census_terms_correct (yes / no)", "reads_naturally (1 poor / 2 acceptable / 3 good)",
             "corrected_question (only if needed)", "corrected_options (only if needed)", "comment"]
PAIR_COLS = ["pair_id", "toggle_axis", "english_A", "english_B", "translated_A", "translated_B",
             "meaning_preserved (yes / partly / no)", "A_B_differ_only_in_identity (yes / no)", "reads_naturally (1 poor / 2 acceptable / 3 good)",
             "corrected_A (only if needed)", "corrected_B (only if needed)", "comment"]


def _sample_facts(rows: List[Dict], n: int, rng: random.Random) -> List[Dict]:
    strata: Dict[tuple, List[Dict]] = {}
    for r in rows:
        strata.setdefault((r["axis"], r["condition"]), []).append(r)
    per = max(1, n // len(strata)); out = []
    for k in sorted(strata):
        g = strata[k][:]; rng.shuffle(g); out += g[:per]
    rest = [r for r in rows if r not in out]; rng.shuffle(rest)
    return (out + rest)[:n]


def _sample_pairs(rows: List[Dict], n: int, rng: random.Random) -> List[Dict]:
    strata: Dict[str, List[Dict]] = {}
    for r in rows:
        strata.setdefault(r["toggle_axis"], []).append(r)
    per = max(1, n // len(strata)); out = []
    for k in sorted(strata):
        g = strata[k][:]; rng.shuffle(g); out += g[:per]
    rest = [r for r in rows if r not in out]; rng.shuffle(rest)
    return (out + rest)[:n]


def _sheet(ws, cols: List[str], rows: List[List[str]], widths: Dict[int, int], validations: Dict[int, List[str]]):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDE8D5"); c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col_idx, options in validations.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{get_column_letter(col_idx)}2:{get_column_letter(col_idx)}{len(rows) + 1}")
    ws.freeze_panes = "A2"


def _workbook(lang: str, facts: List[Dict], pairs: List[Dict], rater: str) -> Workbook:
    wb = Workbook(); ws = wb.active; ws.title = "AgriFacts"
    frows = [[r["id"], r["axis"], r["condition"], r["question_en"], " | ".join(r["choices_en"]), r["question"], " | ".join(r["choices"]),
              "", "", "", "", "", ""] for r in facts]
    _sheet(ws, FACT_COLS, frows, {4: 48, 5: 30, 6: 48, 7: 30, 11: 40, 12: 30, 13: 30}, {8: ["yes", "partly", "no"], 9: ["yes", "no"], 10: ["1", "2", "3"]})
    ws2 = wb.create_sheet("AgriAdvice")
    prows = [[r["pair_id"], r["toggle_axis"], r["version_A_en"]["prompt"], r["version_B_en"]["prompt"], r["version_A"]["prompt"], r["version_B"]["prompt"],
              "", "", "", "", "", ""] for r in pairs]
    _sheet(ws2, PAIR_COLS, prows, {3: 44, 4: 44, 5: 44, 6: 44, 10: 40, 11: 40, 12: 30}, {7: ["yes", "partly", "no"], 8: ["yes", "no"], 9: ["1", "2", "3"]})
    ws3 = wb.create_sheet("Glossary")
    ws3.append(["English term", f"Fixed {LANG_NAME[lang]} rendering (must be used)"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for en, tr in GLOSSARY[lang].items():
        ws3.append([en, tr])
    ws3.column_dimensions["A"].width = 32; ws3.column_dimensions["B"].width = 40
    ws4 = wb.create_sheet("Rater")
    ws4.append(["rater_id", rater]); ws4.append(["language", LANG_NAME[lang]]); ws4.append(["date_completed", ""])
    return wb


def _instructions(n_facts: int, n_pairs: int) -> str:
    return f"""AgriFair — human verification of the Hindi and Bengali translations
=====================================================================

Thank you for checking these translations. The English text is the original. The Hindi and
Bengali text was produced by machine translation and reviewed by automated tools. Before the
files are used in research, a trained human eye must check a sample. Your names will not be
published; ratings are reported in aggregate only.

The package
-----------
  INSTRUCTIONS.txt / INSTRUCTIONS.docx    this document
  AgriFair_Hindi_verification_R1.xlsx     Hindi, for verifier 1
  AgriFair_Hindi_verification_R2.xlsx     Hindi, for verifier 2   (same rows as R1)
  AgriFair_Bengali_verification_R1.xlsx   Bengali, for verifier 1
  AgriFair_Bengali_verification_R2.xlsx   Bengali, for verifier 2 (same rows as R1)

How the two of you should work
------------------------------
  - Decide who is verifier 1 (R1) and who is verifier 2 (R2), and use only your own files.
  - Both of you rate the SAME rows, independently. Please do not discuss rows or compare
    answers until both files are returned; we measure how often two readers agree, and that
    measurement is only meaningful if the ratings are independent.
  - If one of you reads only one of the two languages, say so and rate only that language.

Each workbook has four sheets:
  AgriFacts   {n_facts} census comparison questions with their three answer options.
  AgriAdvice  {n_pairs} pairs of farmer questions. Within a pair, A and B must be identical
              except for the words that say who is asking (for example a woman vs a man).
  Glossary    the fixed wording for census terms. Translations must use these exactly.
  Rater       fill in the date when you finish.

How to rate (about 2 minutes per row; roughly 5 hours per workbook)
-------------------------------------------------------------------
Read the English, then the translation. Fill the empty columns (drop-down lists are provided):

  meaning_preserved             yes = same meaning, nothing added or missing;
                                partly = understandable but a detail is off;
                                no = wrong or misleading.
  census_terms_correct          (AgriFacts) yes if every census term uses the Glossary wording.
  A_B_differ_only_in_identity   (AgriAdvice) yes if A and B are word-for-word identical apart
                                from the identity phrase.
  reads_naturally               1 = a farmer or extension officer would find it awkward or unclear;
                                2 = acceptable; 3 = reads as if written in the language.
  corrected_...                 only if you rated "partly" or "no", or naturalness 1: write the
                                corrected text in full. Keep the option order unchanged. Keep
                                numbers as digits (2015-16, not words).
  comment                       anything else, briefly.

Please do not
-------------
  - change the English columns or the id columns;
  - reorder or delete rows;
  - "improve" translations that are already correct — we need to know how many are right as they are;
  - use machine translation tools while rating.

When done
---------
Fill the Rater sheet (date), save each file under its original name, and send all your files back.
If a row is unclear, leave the rating empty and say why in the comment.

Contact: Koushik Deb, koushik_phd21@iiitkalyani.ac.in
"""


def _docx(work: Path) -> None:
    """INSTRUCTIONS.docx from INSTRUCTIONS.txt via the docx npm package (skipped if node is absent)."""
    import shutil, subprocess
    if not shutil.which("node"):
        return
    script = r"""
const fs=require('fs');const {Document,Packer,Paragraph,TextRun,HeadingLevel}=require('docx');
const lines=fs.readFileSync(process.argv[2],'utf8').split('\n');const ch=[];
for(let i=0;i<lines.length;i++){const l=lines[i];const nxt=lines[i+1]||'';
 if(/^=+$/.test(nxt)){ch.push(new Paragraph({text:l,heading:HeadingLevel.TITLE}));i++;continue;}
 if(/^-+$/.test(nxt)){ch.push(new Paragraph({text:l,heading:HeadingLevel.HEADING_2}));i++;continue;}
 if(/^=+$|^-+$/.test(l))continue;
 const mono=/^ {2}\S/.test(l);ch.push(new Paragraph({spacing:{after:mono?40:120},children:[new TextRun({text:l,font:mono?'Consolas':'Calibri',size:mono?20:22})]}));}
Packer.toBuffer(new Document({sections:[{children:ch}]})).then(b=>fs.writeFileSync(process.argv[3],b));
"""
    js = work / "_mk.js"; js.write_text(script)
    env = dict(**__import__("os").environ); env["NODE_PATH"] = str(Path(__file__).resolve().parents[3] / "node_modules") + ":" + env.get("NODE_PATH", "")
    r = subprocess.run(["node", str(js), str(work / "INSTRUCTIONS.txt"), str(work / "INSTRUCTIONS.docx")], capture_output=True, text=True, env=env)
    js.unlink()
    if r.returncode != 0:
        print("docx skipped:", r.stderr.strip()[:120])


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(); ap.add_argument("--facts", type=int, default=100); ap.add_argument("--pairs", type=int, default=50)
    ap.add_argument("--seed", type=int, default=20260921); a = ap.parse_args(argv)
    DEST.mkdir(parents=True, exist_ok=True); rng = random.Random(a.seed)
    work = DEST / "package"; work.mkdir(exist_ok=True)
    manifest = {"seed": a.seed, "facts_per_language": a.facts, "pairs_per_language": a.pairs, "raters": RATERS, "languages": {}}
    files = []
    for lang in CFG["languages"]:
        facts = _sample_facts(_jsonl(OUT / f"agrifacts_{lang}.jsonl"), a.facts, rng)
        pairs = _sample_pairs(_jsonl(OUT / f"agriadvice_{lang}.jsonl"), a.pairs, rng)
        for rater in RATERS:
            name = f"AgriFair_{LANG_NAME[lang]}_verification_{rater}.xlsx"
            _workbook(lang, facts, pairs, rater).save(work / name); files.append(name)
        manifest["languages"][lang] = {"fact_ids": [r["id"] for r in facts], "pair_ids": [r["pair_id"] for r in pairs],
                                       "fact_sha256": {r["id"]: hashlib.sha256((r["question"] + "|" + "|".join(r["choices"])).encode()).hexdigest() for r in facts},
                                       "pair_sha256": {r["pair_id"]: hashlib.sha256((r["version_A"]["prompt"] + "|" + r["version_B"]["prompt"]).encode()).hexdigest() for r in pairs}}
        print(f"{lang}: {len(facts)} facts, {len(pairs)} pairs")
    (work / "INSTRUCTIONS.txt").write_text(_instructions(a.facts, a.pairs), encoding="utf-8")
    _docx(work); files.append("INSTRUCTIONS.txt")
    if (work / "INSTRUCTIONS.docx").exists():
        files.append("INSTRUCTIONS.docx")
    zpath = DEST / "AgriFair_translation_verification.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(work / f, f)
    (DEST / "verification_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    print("zip ->", zpath, "| manifest ->", DEST / "verification_manifest.json")


if __name__ == "__main__":
    main()
