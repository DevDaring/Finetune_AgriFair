"""Import the returned R3 advice ratings (pack 2), research content only.

    python -m Submission1_Code_Phase2.import_r3_ratings \
        --files Submission1/phase2_rater_pack/AgriFair_advice_ratings_R1.xlsx \
                Submission1/phase2_rater_pack/AgriFair_advice_ratings_R2.xlsx

What is read: the scores, the change judgements, and free text that describes the *task*.

What is never read: any name, initials, employer, contact or signature cell, and any column a
reader may have added whose header names a person. Readers are "R1" and "R2" by filename order
only; nothing in the outputs identifies who did the work.

Free-text comments are scrubbed before they enter the results tree. The allowlist of acronyms to
KEEP is derived from our own source material (the reference packets and the prompts), not typed by
hand and not read off the returned sheets: a capitalised token that occurs in material we wrote is
a domain term, and any other capitalised token is treated as possible initials and replaced. That
rule means no one has to read the returned free text to decide what is safe to keep.
"""
from __future__ import annotations

import argparse
import collections
import csv
import json
import re
from pathlib import Path
from typing import Dict, List

from openpyxl import load_workbook

from Submission1_Code_Phase2 import common as C
from Submission1_Code_Phase2.import_rater_protocol import (
    IDENTITY_COLUMNS, IDENTITY_PATTERNS, _KEEP_ACRONYMS, _INITIALS)

ENTRY_NUMERIC = ["left_essentials_covered_count", "right_essentials_covered_count",
                 "left_unsupported_or_incorrect_claims", "right_unsupported_or_incorrect_claims"]
ENTRY_ENUM = {"left_unsafe_yes_no": {"yes", "no"}, "right_unsafe_yes_no": {"yes", "no"},
              "left_usefulness_0_2": {"0", "1", "2"}, "right_usefulness_0_2": {"0", "1", "2"},
              "pair_change_type": {"none", "presentation_only", "substantive"},
              "change_is_justified_yes_no_na": {"yes", "no", "not_applicable"}}


def domain_acronyms(advice: Path) -> set:
    """Acronyms that occur in material WE wrote. Everything else is treated as possible initials."""
    text = []
    p = advice / "reference_packets.csv"
    if p.exists():
        text.append(p.read_text(encoding="utf-8", errors="ignore"))
    q = advice / "r3_prompts.jsonl"
    if q.exists():
        text.append(q.read_text(encoding="utf-8", errors="ignore"))
    found = set(_INITIALS.findall("\n".join(text)))
    return _KEEP_ACRONYMS | found


# Dotted initials, of the form "X.Y." or "X.Y.Z". The consecutive-caps pattern in
# import_rater_protocol does NOT match these -- the letters are separated by periods -- so every
# signed comment would otherwise pass straight through into the results tree.
_DOTTED = re.compile(r"\b(?:[A-Z]\.){1,3}(?:[A-Z]\b)?")


def scrub(text: str, keep: set) -> str:
    if not text:
        return ""
    # A dotted form is a domain term only if its letters spell one that occurs in our own
    # material (so a fertiliser grade survives); anything else is treated as initials.
    out = _DOTTED.sub(
        lambda m: m.group(0) if m.group(0).replace(".", "") in keep else "[reviewer]", text)
    out = _INITIALS.sub(lambda m: m.group(0) if m.group(0) in keep else "[reviewer]", out)
    out = re.sub(r"(\[reviewer\]\W*){2,}", "[reviewer] ", out)
    # drop anything that looks like contact detail regardless of case
    out = re.sub(r"[\w.+-]+@[\w-]+\.\w+", "[contact removed]", out)
    out = re.sub(r"\b(?:\+91[\s-]?)?[6-9]\d{9}\b", "[contact removed]", out)
    out = re.sub(r"\b(?:reviewed|checked|rated|signed)\s+by\b.*?(?=[.;]|$)", "[attribution removed]",
                 out, flags=re.I)
    return re.sub(r"^\s*\[reviewer\]\s*[.,:;-]*\s*", "", out).strip()


def read_ratings(path: Path, rater: str, keep: set) -> List[Dict]:
    wb = load_workbook(path, data_only=True)
    sheet = next((s for s in wb.sheetnames if s.lower().startswith("ratings")), wb.sheetnames[0])
    ws = wb[sheet]
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    drop = {i for i, h in enumerate(hdr)
            if not h or h.lower() in IDENTITY_COLUMNS or IDENTITY_PATTERNS.search(h)}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v in (None, "") for v in r):
            continue
        d = {}
        for i, h in enumerate(hdr):
            if i in drop or i >= len(r):
                continue
            v = "" if r[i] is None else str(r[i]).strip()
            d[h] = scrub(v, keep) if h == "comment" else v
        d["rater"] = rater
        rows.append(d)
    return rows, sorted(hdr[i] for i in drop if hdr[i])


def validate(rows: List[Dict], rater: str) -> Dict:
    """Value-domain and range checks. Blanks are legal where the rubric allows them."""
    problems = collections.Counter()
    for r in rows:
        total = r.get("essentials_total", "")
        total = int(float(total)) if str(total).strip() else 0
        for col in ENTRY_NUMERIC:
            v = str(r.get(col, "")).strip()
            if not v:
                continue
            try:
                n = float(v)
            except ValueError:
                problems[f"{col}:not_a_number"] += 1
                continue
            if n < 0:
                problems[f"{col}:negative"] += 1
            if col.endswith("essentials_covered_count") and total and n > total:
                problems[f"{col}:above_denominator"] += 1
        for col, allowed in ENTRY_ENUM.items():
            v = str(r.get(col, "")).strip()
            if v and v not in allowed:
                problems[f"{col}:unexpected_value"] += 1
        # rubric rule 6: not_applicable exactly when the change is none/presentation_only
        ch, ju = r.get("pair_change_type", ""), r.get("change_is_justified_yes_no_na", "")
        if ch in ("none", "presentation_only") and ju and ju != "not_applicable":
            problems["rule6:judged_a_non_substantive_change"] += 1
        if ch == "substantive" and ju == "not_applicable":
            problems["rule6:na_on_a_substantive_change"] += 1
    return {"rater": rater, "rows": len(rows), "problems": dict(problems)}


def main(argv=None) -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--files", nargs="+", type=Path, required=True)
    a = ap.parse_args(argv)
    cfg = C.load_config()
    out = C.CODES_ROOT / cfg["output_directory"] / "advice"
    keep = domain_acronyms(out)

    key = {r["assessment_id"]: r for r in
           csv.DictReader((out / "KEEP_FROM_RATERS_r3_key.csv").open(encoding="utf-8"))}

    report, all_rows, dropped_cols = [], [], set()
    for i, f in enumerate(sorted(a.files), start=1):
        rater = f"R{i}"
        rows, dropped = read_ratings(f, rater, keep)
        dropped_cols |= set(dropped)
        unknown = [r["assessment_id"] for r in rows if r.get("assessment_id") not in key]
        rep = validate(rows, rater)
        rep["unknown_assessment_ids"] = len(unknown)
        rep["source_file"] = f.name
        report.append(rep)
        all_rows += rows

    C.write_csv(out / "r3_returned_ratings.csv", all_rows)
    # per-rater scrubbed copies, so the downstream analyser never touches the returned workbooks
    for rep in report:
        rid = rep["rater"]
        C.write_csv(out / f"r3_ratings_scrubbed_{rid}.csv", [r for r in all_rows if r["rater"] == rid])
    summary = {"raters": len(a.files), "rows_total": len(all_rows),
               "identity_columns_dropped": sorted(dropped_cols),
               "acronyms_kept_from_our_own_material": len(keep),
               "per_rater": report,
               "note": ("Readers are R1/R2 by filename order. No name, initials, employer or contact "
                        "value is read from the returned files; free text is scrubbed on the way in.")}
    C.write_json(out / "r3_import_report.json", summary)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
